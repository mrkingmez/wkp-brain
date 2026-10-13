# -*- coding: utf-8 -*-
"""One-off markdown -> docx converter, scratch utility.
Handles: #/##/### headers, **bold**, *italics*, `code`, > blockquotes,
- bullets, --- rules, and plain paragraphs. Not a full CommonMark parser.
Usage: python md_to_docx.py <input.md> <output.docx>
"""
import sys
import re
from docx import Document
from docx.shared import Pt

def add_runs(paragraph, text):
    # split on **bold**, *italic*, `code` while preserving order
    tokens = re.split(r'(\*\*.+?\*\*|\*.+?\*|`.+?`)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            r = paragraph.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith('`') and tok.endswith('`'):
            r = paragraph.add_run(tok[1:-1])
            r.font.name = 'Consolas'
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            r = paragraph.add_run(tok[1:-1])
            r.italic = True
        else:
            paragraph.add_run(tok)

def main():
    src, dst = sys.argv[1], sys.argv[2]
    with open(src, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()
    table_rows = []

    def flush_table():
        if not table_rows:
            return
        ncols = max(len(r) for r in table_rows)
        t = doc.add_table(rows=0, cols=ncols)
        t.style = 'Table Grid'
        for ri, r in enumerate(table_rows):
            cells = t.add_row().cells
            for ci in range(ncols):
                para = cells[ci].paragraphs[0]
                add_runs(para, r[ci] if ci < len(r) else '')
                if ri == 0:
                    for run in para.runs:
                        run.bold = True
        table_rows.clear()

    for raw in lines:
        line = raw.rstrip('\n')
        stripped = line.strip()

        if stripped.startswith('|'):
            cells = [c.strip() for c in stripped.strip('|').split('|')]
            if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells):
                table_rows.append(cells)
            continue
        flush_table()

        if not stripped:
            continue
        if stripped == '---':
            doc.add_paragraph('_' * 60)
            continue
        if stripped.startswith('### '):
            doc.add_heading(stripped[4:], level=3)
            continue
        if stripped.startswith('## '):
            doc.add_heading(stripped[3:], level=2)
            continue
        if stripped.startswith('# '):
            doc.add_heading(stripped[2:], level=1)
            continue
        if stripped.startswith('> '):
            p = doc.add_paragraph(style='Intense Quote')
            add_runs(p, stripped[2:])
            continue
        if stripped == '>':
            continue
        if stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, stripped[2:])
            continue
        if re.match(r'^\d+\.\s', stripped):
            p = doc.add_paragraph(style='List Number')
            add_runs(p, re.sub(r'^\d+\.\s', '', stripped))
            continue

        p = doc.add_paragraph()
        add_runs(p, stripped)

    flush_table()
    doc.save(dst)
    print(f"Wrote {dst}")

if __name__ == '__main__':
    main()
