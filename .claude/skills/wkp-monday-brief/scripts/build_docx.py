# -*- coding: utf-8 -*-
"""Markdown -> docx converter for the WKP Monday Brief, with real table support.
Handles: #/##/### headers, **bold**, *italics*, `code`, > blockquotes,
- bullets, 1. numbered lists, --- rules, plain paragraphs, and GFM pipe tables.
Usage: python build_docx.py <input.md> <output.docx>
"""
import sys
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_runs(paragraph, text):
    tokens = re.split(r'(\*\*.+?\*\*|\*.+?\*|`.+?`)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            r = paragraph.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith('`') and tok.endswith('`'):
            r = paragraph.add_run(tok[1:-1])
            r.font.name = 'Consolas'
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            r = paragraph.add_run(tok[1:-1])
            r.italic = True
        else:
            paragraph.add_run(tok)


def is_table_row(line):
    return line.strip().startswith('|') and line.strip().endswith('|')


def is_separator_row(line):
    cells = [c.strip() for c in line.strip().strip('|').split('|')]
    return all(re.match(r'^:?-+:?$', c) for c in cells if c != '')


def parse_row(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]


def main():
    src, dst = sys.argv[1], sys.argv[2]
    with open(src, 'r', encoding='utf-8') as f:
        lines = [l.rstrip('\n') for l in f.readlines()]

    doc = Document()
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if is_table_row(stripped) and i + 1 < n and is_separator_row(lines[i + 1]):
            header = parse_row(stripped)
            i += 2
            rows = []
            while i < n and is_table_row(lines[i].strip()):
                rows.append(parse_row(lines[i].strip()))
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = 'Light Grid Accent 1'
            for c, val in enumerate(header):
                cell_p = table.rows[0].cells[c].paragraphs[0]
                r = cell_p.add_run(val)
                r.bold = True
            for row in rows:
                cells = table.add_row().cells
                for c, val in enumerate(row):
                    if c < len(cells):
                        add_runs(cells[c].paragraphs[0], val)
            doc.add_paragraph('')
            continue

        if stripped == '---':
            doc.add_paragraph('_' * 60)
            i += 1
            continue
        if stripped.startswith('### '):
            doc.add_heading(stripped[4:], level=3)
            i += 1
            continue
        if stripped.startswith('## '):
            doc.add_heading(stripped[3:], level=2)
            i += 1
            continue
        if stripped.startswith('# '):
            h = doc.add_heading(stripped[2:], level=1)
            i += 1
            continue
        if stripped.startswith('> '):
            p = doc.add_paragraph(style='Intense Quote')
            add_runs(p, stripped[2:])
            i += 1
            continue
        if stripped == '>':
            i += 1
            continue
        if stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, stripped[2:])
            i += 1
            continue
        if re.match(r'^\d+\.\s', stripped):
            p = doc.add_paragraph(style='List Number')
            add_runs(p, re.sub(r'^\d+\.\s', '', stripped))
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    doc.save(dst)
    print(f"Wrote {dst}")


if __name__ == '__main__':
    main()
