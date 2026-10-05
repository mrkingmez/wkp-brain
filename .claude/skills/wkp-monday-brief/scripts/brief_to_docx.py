# -*- coding: utf-8 -*-
"""Markdown -> docx for the Monday brief. Handles #/##/###, **bold**, *italic*,
`code`, - and 1. lists, --- rules, and pipe tables.
Usage: python brief_to_docx.py <input.md> <output.docx>"""
import sys, re
from docx import Document
from docx.shared import Pt

def runs(p, text):
    for tok in re.split(r'(\*\*.+?\*\*|`.+?`|\*[^*\s][^*]*?\*)', text):
        if not tok: continue
        if tok.startswith('**'): p.add_run(tok[2:-2]).bold = True
        elif tok.startswith('`'): p.add_run(tok[1:-1]).font.name = 'Consolas'
        elif tok.startswith('*') and len(tok) > 2: p.add_run(tok[1:-1]).italic = True
        else: p.add_run(tok)

def main(src, dst):
    d = Document()
    d.styles['Normal'].font.name = 'Calibri'
    d.styles['Normal'].font.size = Pt(11)
    lines = open(src, encoding='utf-8').read().split('\n')
    i = 0
    while i < len(lines):
        l = lines[i].rstrip()
        if l.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', c) for c in cells): rows.append(cells)
                i += 1
            t = d.add_table(rows=len(rows), cols=len(rows[0]), style='Table Grid')
            for r, row in enumerate(rows):
                for c, txt in enumerate(row):
                    cell = t.cell(r, c); cell.text = ''
                    runs(cell.paragraphs[0], txt)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(9.5)
                        if r == 0: run.bold = True
            d.add_paragraph()
            continue
        if l.startswith('### '): d.add_heading(l[4:], 3)
        elif l.startswith('## '): d.add_heading(l[3:], 2)
        elif l.startswith('# '): d.add_heading(l[2:], 1)
        elif l.strip() == '---': pass
        elif re.match(r'^\d+\. ', l): runs(d.add_paragraph(style='List Number'), re.sub(r'^\d+\. ', '', l))
        elif l.startswith('- '): runs(d.add_paragraph(style='List Bullet'), l[2:])
        elif l.strip(): runs(d.add_paragraph(), l)
        i += 1
    d.save(dst)

main(sys.argv[1], sys.argv[2])
